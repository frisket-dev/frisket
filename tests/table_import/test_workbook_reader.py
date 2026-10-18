from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import shutil

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from frisket.table_import import (
    TableReadSettings,
    count_records,
    inspect_source,
    open_table,
)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Visible"
    sheet.append(["name", "amount", "when", "active"])
    sheet.append(["Ada", 1.25, date(2026, 10, 18), True])
    sheet.append(["Grace", 2, datetime(2026, 10, 19, 12, 30), False])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["only"])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart_sheet = workbook.create_chartsheet("Chart only")
    chart_sheet.add_chart(chart)
    workbook.save(path)


def test_xlsx_inspection_and_iteration_keep_typed_semantics(tmp_path: Path) -> None:
    source = tmp_path / "book.data"
    _workbook(source)

    info = inspect_source(source)
    assert info.kind == "xlsx"
    assert [
        (sheet.index, sheet.name, sheet.hidden, sheet.empty) for sheet in info.sheets
    ] == [
        (0, "Visible", False, False),
        (1, "Hidden", True, False),
        (2, "Chart only", False, True),
    ]
    settings = TableReadSettings(kind="xlsx", worksheet="Visible")
    assert count_records(source, settings) == 3

    with open_table(source, settings, total_record_count=3) as reader:
        rows = list(reader)

    assert [column.name for column in reader.columns] == [
        "name",
        "amount",
        "when",
        "active",
    ]
    assert [cell.kind for cell in rows[1].cells] == [
        "text",
        "number",
        "date",
        "boolean",
    ]
    assert [cell.text for cell in rows[1].cells] == [
        "Ada",
        "1.25",
        "2026-10-18",
        "true",
    ]
    assert rows[2].cells[2].kind == "datetime"
    assert rows[2].cells[2].text == "2026-10-19T12:30:00"


def test_calamine_reads_extensionless_xls_and_ods_sources(tmp_path: Path) -> None:
    fixtures = Path(__file__).with_name("data")
    for fixture_name, expected_kind in (("base.xls", "xls"), ("base.ods", "ods")):
        source = tmp_path / f"staged-{expected_kind}"
        shutil.copyfile(fixtures / fixture_name, source)

        info = inspect_source(source)
        assert info.kind == expected_kind
        assert info.sheets
        settings = TableReadSettings(kind=expected_kind, worksheet=info.sheets[0].name)  # type: ignore[arg-type]
        total = count_records(source, settings)
        assert total > 1
        with open_table(source, settings, total_record_count=total) as reader:
            first_rows = []
            for index, row in enumerate(reader):
                first_rows.append(row)
                if index == 1:
                    break
        assert first_rows[0].role == "header"
        assert first_rows[1].cells

from __future__ import annotations

import csv
import io
import json
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _seed(client: TestClient) -> tuple[str, int, int]:
    pid = client.post("/api/projects", json={"name": "Civic résumé"}).json()["id"]
    project = client.app.state.workspace.get(pid)

    cases_id = project.add_sheet("Cases / 2026")
    case_columns = {
        "person": project.add_column(cases_id, "person"),
        "status": project.add_column(cases_id, "status"),
        "amount": project.add_column(cases_id, "amount", type="number"),
    }
    project.add_rows(
        cases_id,
        [
            {"person": "José", "status": "open", "amount": 12},
            {"person": "李雷", "status": "closed", "amount": -3.5},
        ],
        case_columns,
    )

    notes_id = project.add_sheet("Notes: 2026")
    note_columns = {
        "speaker": project.add_column(notes_id, "speaker"),
        "quote": project.add_column(notes_id, "quote"),
    }
    project.add_rows(
        notes_id,
        [{"speaker": "Zoë", "quote": "“smart quotes”"}],
        note_columns,
    )
    return pid, cases_id, notes_id


def test_selected_sheets_export_as_unicode_excel_tabs(
    replay_client: TestClient,
) -> None:
    pid, cases_id, notes_id = _seed(replay_client)
    response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params=[("sheet_id", cases_id), ("sheet_id", notes_id), ("format", "xlsx")],
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="Civic-r-sum.xlsx"' in response.headers["content-disposition"]

    workbook = load_workbook(
        io.BytesIO(response.content), read_only=True, data_only=False
    )
    assert workbook.sheetnames == ["Cases - 2026", "Notes- 2026"]
    assert list(workbook["Cases - 2026"].values) == [
        ("person", "status", "amount"),
        ("José", "open", 12),
        ("李雷", "closed", -3.5),
    ]
    assert list(workbook["Notes- 2026"].values) == [
        ("speaker", "quote"),
        ("Zoë", "“smart quotes”"),
    ]


def test_selected_sheets_export_as_bom_prefixed_csv_zip(
    replay_client: TestClient,
) -> None:
    pid, cases_id, notes_id = _seed(replay_client)
    response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params=[("sheet_id", cases_id), ("sheet_id", notes_id), ("format", "csv")],
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/zip")
    assert 'filename="Civic-r-sum-csv.zip"' in response.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == ["Cases-2026.csv", "Notes-2026.csv"]
        cases_bytes = archive.read("Cases-2026.csv")
        notes_bytes = archive.read("Notes-2026.csv")
    assert cases_bytes.startswith(b"\xef\xbb\xbf")
    assert notes_bytes.startswith(b"\xef\xbb\xbf")
    assert list(csv.reader(io.StringIO(cases_bytes.decode("utf-8-sig")))) == [
        ["person", "status", "amount"],
        ["José", "open", "12"],
        ["李雷", "closed", "-3.5"],
    ]
    assert "Zoë" in notes_bytes.decode("utf-8-sig")


def test_single_sheet_xlsx_can_match_the_filtered_current_view(
    replay_client: TestClient,
) -> None:
    pid, cases_id, _notes_id = _seed(replay_client)
    response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params=[
            ("sheet_id", cases_id),
            ("format", "xlsx"),
            ("filter", json.dumps({"status": {"eq": "open"}})),
        ],
    )

    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    assert list(workbook["Cases - 2026"].values) == [
        ("person", "status", "amount"),
        ("José", "open", 12),
    ]


def test_direct_csv_has_excel_friendly_utf8_bom(replay_client: TestClient) -> None:
    pid, cases_id, _notes_id = _seed(replay_client)
    response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params={"sheet_id": cases_id, "format": "csv"},
    )

    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert response.encoding == "utf-8-sig"
    assert "José" in response.text


def test_multi_sheet_current_view_is_rejected(replay_client: TestClient) -> None:
    pid, cases_id, notes_id = _seed(replay_client)
    response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params=[
            ("sheet_id", cases_id),
            ("sheet_id", notes_id),
            ("format", "xlsx"),
            ("filter", json.dumps({"status": {"eq": "open"}})),
        ],
    )

    assert response.status_code == 400
    assert response.json() == {
        "detail": "current-view export requires exactly one selected sheet"
    }


def test_formula_like_headers_are_escaped_in_csv_zip_and_excel(
    replay_client: TestClient,
) -> None:
    pid = replay_client.post("/api/projects", json={"name": "Formula headers"}).json()[
        "id"
    ]
    project = replay_client.app.state.workspace.get(pid)
    dangerous_header = '=HYPERLINK("https://example.test","click")'
    dangerous_sheet_id = project.add_sheet("Dangerous")
    dangerous_column_id = project.add_column(dangerous_sheet_id, dangerous_header)
    project.add_rows(
        dangerous_sheet_id,
        [{dangerous_header: "ordinary value"}],
        {dangerous_header: dangerous_column_id},
    )
    companion_sheet_id = project.add_sheet("Companion")
    companion_column_id = project.add_column(companion_sheet_id, "safe")
    project.add_rows(
        companion_sheet_id,
        [{"safe": "value"}],
        {"safe": companion_column_id},
    )

    csv_response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params={"sheet_id": dangerous_sheet_id, "format": "csv"},
    )
    assert csv_response.status_code == 200, csv_response.text
    assert next(csv.reader(io.StringIO(csv_response.text)))[0] == f"'{dangerous_header}"

    zip_response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params=[
            ("sheet_id", dangerous_sheet_id),
            ("sheet_id", companion_sheet_id),
            ("format", "csv"),
        ],
    )
    assert zip_response.status_code == 200, zip_response.text
    with zipfile.ZipFile(io.BytesIO(zip_response.content)) as archive:
        header = next(
            csv.reader(io.StringIO(archive.read("Dangerous.csv").decode("utf-8-sig")))
        )[0]
    assert header == f"'{dangerous_header}"

    xlsx_response = replay_client.get(
        f"/api/projects/{pid}/exports/sheets",
        params={"sheet_id": dangerous_sheet_id, "format": "xlsx"},
    )
    assert xlsx_response.status_code == 200, xlsx_response.text
    workbook = load_workbook(
        io.BytesIO(xlsx_response.content), read_only=True, data_only=False
    )
    assert workbook["Dangerous"]["A1"].value == f"'{dangerous_header}"
    assert workbook["Dangerous"]["A1"].data_type == "s"

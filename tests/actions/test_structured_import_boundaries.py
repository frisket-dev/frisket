from __future__ import annotations

import hashlib
import json
from contextlib import closing

import pytest
from openpyxl import Workbook
from pydantic import BaseModel

from frisket.actions.core import ActionCategory, ActionNamespace, ActionRegistry
from frisket.actions.imports import FileSource
from frisket.actions.system import BoundTypedActionRequest
from frisket.actions.types import ActionRequest, LocalFileReader
from frisket.engine.executor import ExecutorDeps, ImportWorkloadLimits, run_action_spec
from frisket.engine.executor.table_action import run_typed_create_sheet_action
from frisket.engine.store import Project
from frisket.engine.store.receipts import ReceiptStore
from frisket.sdk import ActionParams, TableResult, TableRow, action, create_sheet


def _request(tmp_path, kind):
    path = tmp_path / kind
    if kind == "import.xlsx":
        workbook = Workbook()
        workbook.active.append(["name"])
        workbook.active.append(["Ada"])
        workbook.active.append(["Grace"])
        workbook.save(path)
        workbook.close()
        options = {"columns": [{"name": "name", "type": "text"}]}
    elif kind == "import.geojson":
        path.write_text(
            json.dumps(
                {
                    "type": "FeatureCollection",
                    "features": [
                        {
                            "type": "Feature",
                            "geometry": {"type": "Point", "coordinates": [n, 0]},
                            "properties": {},
                        }
                        for n in (1, 2)
                    ],
                }
            )
        )
        options = {}
    else:
        path.write_text(
            "<kml><Document><Placemark><Point><coordinates>1,0</coordinates></Point></Placemark><Placemark><Point><coordinates>2,0</coordinates></Point></Placemark></Document></kml>"
        )
        options = {}
    return path, {
        "action_id": kind,
        "scope": {"kind": "project"},
        "sheet_name": "Imported",
        "params": {"source": {"kind": "file", "path": str(path)}, **options},
        "idempotency_key": "structured-import",
    }


@pytest.mark.parametrize("kind", ["import.xlsx", "import.geojson", "import.kml"])
def test_typed_import_budget_replay_and_retired_envelope_refusal(tmp_path, kind):
    path, request = _request(tmp_path, kind)
    with closing(Project.create(tmp_path / "project")) as project:
        refused = run_action_spec(
            project,
            request,
            project_id="p",
            deps=ExecutorDeps(import_workload_limits=ImportWorkloadLimits(max_rows=1)),
        )
        assert refused.status == "failed"
        assert refused.errors[0].code == "import_workload_limit_exceeded"
        assert project.db.execute("SELECT COUNT(*) FROM sheets").fetchone()[0] == 0
        completed = run_action_spec(project, request, project_id="p")
        assert completed.status == "completed", completed.errors
        path.unlink()
        replay = run_action_spec(project, request, project_id="p")
        assert replay.receipt_id == completed.receipt_id
        legacy = {
            "schema_version": "frisket.action.v2",
            "kind": kind,
            "capabilities": ["project:write"],
            "params": {
                "sheet_name": "Imported",
                "mode": "create_sheet",
                **request["params"],
            },
            "idempotency_key": request["idempotency_key"],
        }
        refused = run_action_spec(project, legacy, project_id="p")
        assert refused.status == "failed"
        assert refused.receipt_id is None
        assert project.db.execute("SELECT COUNT(*) FROM sheets").fetchone()[0] == 1
        assert project.db.execute("SELECT COUNT(*) FROM receipts").fetchone()[0] == 1


class SourceParams(ActionParams):
    source: FileSource


@pytest.mark.parametrize(
    ("header", "columns", "accepted"),
    [
        (["b", "a", "secret"], [{"name": "a"}], False),
        (["b", "a"], [{"name": "a"}, {"name": "b"}], False),
        (["a", "a"], [{"name": "a", "source_name": "a"}], False),
        (["b", "a", "secret"], [{"name": "renamed", "source_name": "a"}], True),
    ],
)
def test_xlsx_header_projection_requires_explicit_unambiguous_mapping(
    tmp_path, header, columns, accepted
):
    path = tmp_path / "mapped.xlsx"
    workbook = Workbook()
    workbook.active.append(header)
    workbook.active.append([f"value-{index}" for index in range(len(header))])
    workbook.save(path)
    workbook.close()
    with closing(Project.create(tmp_path / "project")) as project:
        result = run_action_spec(
            project,
            {
                "action_id": "import.xlsx",
                "scope": {"kind": "project"},
                "sheet_name": "Mapped",
                "params": {
                    "source": {"kind": "file", "path": str(path)},
                    "columns": [{"type": "text", **column} for column in columns],
                },
                "idempotency_key": "header-check",
            },
            project_id="p",
        )
        if accepted:
            assert result.status == "completed", result.errors
            sheet = project.sheets()[0]
            column = project.columns(sheet["id"])[0]
            assert project.get_values(sheet["id"], column["id"]) == {1: "value-1"}
        else:
            assert result.status == "failed"
            assert result.errors[0].code == "xlsx_header_mismatch"
            assert project.sheets() == []


class TextRow(BaseModel):
    value: str


def test_file_facts_override_producer_claims_and_keep_quarantine_on_replay(tmp_path):
    path = tmp_path / "actual.txt"
    raw = b"actual bytes"
    path.write_bytes(raw)
    calls = []

    def produce(params: SourceParams, files: LocalFileReader) -> TableResult[TextRow]:
        calls.append(True)
        value = files.read_bytes(params.source.path).decode()
        return TableResult(
            rows=[TableRow(output=TextRow(value=value))],
            source={
                "kind": "file",
                "label": "descriptive label",
                "importer": "custom",
                "path": "forged",
                "fingerprint": "forged",
                "line_count": 999,
                "byte_count": 999,
                "skipped_features": [0],
            },
        )

    registry = ActionRegistry(
        (
            ActionNamespace(
                "custom",
                actions=(
                    action(
                        name="source",
                        title="Source",
                        description="Read a source.",
                        category=ActionCategory.CONVERT,
                        run=create_sheet(produce),
                    ),
                ),
            ),
        )
    )
    request = ActionRequest(
        action_id="custom.source",
        scope={"kind": "project"},
        sheet_name="Read",
        params={"source": {"kind": "file", "path": str(path)}},
        idempotency_key="observed-source",
    )
    bound = BoundTypedActionRequest.bind(registry.get("custom.source"), request)
    with closing(Project.create(tmp_path / "project")) as project:
        result = run_typed_create_sheet_action(project, "p", bound)
        assert result.status == "completed", result.errors
        receipt = ReceiptStore(project).parsed_by_id(result.receipt_id)
        source = next(
            item.ref for item in receipt.inputs if item.ref.get("kind") == "file_rows"
        )
        assert source["path"] == str(path)
        assert source["fingerprint"] == "sha256:" + hashlib.sha256(raw).hexdigest()
        assert source["line_count"] == 1
        assert source["skipped_features"] == [0]
        assert source["label"] == "descriptive label"
        facts = [
            item.ref
            for item in receipt.inputs
            if item.ref.get("kind") == "local_file_read"
        ]
        assert facts == [
            {
                "kind": "local_file_read",
                "path": str(path),
                "sha256": source["fingerprint"],
                "byte_count": len(raw),
            }
        ]
        path.unlink()
        replay = run_typed_create_sheet_action(project, "p", bound)
        assert replay.receipt_id == result.receipt_id
        assert calls == [True]
        assert (
            ReceiptStore(project).parsed_by_id(replay.receipt_id).inputs
            == receipt.inputs
        )

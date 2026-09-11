from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

from executor_harness import CatalogEntry, ExecutorCase, Gate
from frisket.engine.store import Project
from helpers import run_cli
from action_test_helpers import write_json

CSV_BODY = (
    "summary,source_url,score\n"
    "City hall awarded a no-bid contract,https://example.com/story/1,9\n"
    "Routine road work finished early,https://example.com/story/2,2\n"
)


def _write_csv(dir_path: Path, name: str, body: str) -> Path:
    path = dir_path / name
    path.write_text(body, encoding="utf-8")
    return path


def _import_csv_action(
    csv_path: Path,
    *,
    sheet_name: str = "Articles",
    idempotency_key: str | None = "import_csv@sha256:stable",
) -> dict[str, Any]:
    return {
        "action_id": "import.csv",
        "scope": {"kind": "project"},
        "sheet_name": sheet_name,
        "params": {
            "sources": [
                {
                    "path": str(csv_path),
                    "label": "city-contracts.csv",
                    "delimiter": ",",
                    "encoding": "utf-8",
                }
            ],
            "columns": [
                {"name": "summary", "type": "text"},
                {"name": "source_url", "type": "url"},
                {"name": "score", "type": "integer"},
            ],
        },
        "idempotency_key": idempotency_key,
    }


def _seed(project: Project, tmp_path: Path) -> dict[str, Any]:
    del project
    return {"dir": tmp_path, "csv_path": _write_csv(tmp_path, "rows.csv", CSV_BODY)}


def _make_action(seeded: dict[str, Any]) -> dict[str, Any]:
    return _import_csv_action(seeded["csv_path"])


def _missing_idempotency_action(seeded: dict[str, Any]) -> dict[str, Any]:
    return _import_csv_action(seeded["csv_path"], idempotency_key=None)


def _missing_source_action(seeded: dict[str, Any]) -> dict[str, Any]:
    action = _import_csv_action(
        seeded["dir"] / "missing.csv",
        sheet_name="Missing source",
        idempotency_key="import_csv@sha256:missing-source",
    )
    return action


def _header_mismatch_action(seeded: dict[str, Any]) -> dict[str, Any]:
    action = _import_csv_action(
        seeded["csv_path"], idempotency_key="import_csv@sha256:mismatch"
    )
    action["params"]["columns"][1]["name"] = "url"
    return action


def _invalid_value_action(seeded: dict[str, Any]) -> dict[str, Any]:
    csv_path = _write_csv(
        seeded["dir"],
        "invalid.csv",
        "summary,source_url,score\nBad row,https://example.com/b,not-an-int\n",
    )
    return _import_csv_action(
        csv_path, sheet_name="Bad Rows", idempotency_key="import_csv@sha256:invalid"
    )


def _check_state(project: Project, seeded: dict[str, Any], result: Any) -> None:
    from frisket.contracts.action import Receipt

    assert result.op_ids == [1]
    assert [output.kind for output in result.outputs] == [
        "sheet",
        "column",
        "column",
        "column",
        "rows",
    ]

    sheet = project.db.execute("SELECT * FROM sheets WHERE name='Articles'").fetchone()
    assert sheet is not None
    assert project.row_count(int(sheet["id"])) == 2
    columns = {
        row["name"]: row
        for row in project.db.execute(
            "SELECT * FROM columns WHERE sheet_id=? ORDER BY position",
            (sheet["id"],),
        ).fetchall()
    }
    assert columns["source_url"]["type"] == "link"
    assert columns["score"]["type"] == "integer"
    scores = project.get_values(int(sheet["id"]), int(columns["score"]["id"]))
    assert list(scores.values()) == [9, 2]

    op = project.db.execute("SELECT * FROM ops WHERE id=1").fetchone()
    assert op is not None
    assert op["kind"] == "import.csv"
    op_spec = json.loads(op["spec"])
    assert op_spec["action_id"] == "import.csv"
    assert op_spec["scope"] == {"kind": "project"}
    assert op_spec["sheet_name"] == "Articles"
    assert op_spec["params"]["sources"][0]["path"] == str(seeded["csv_path"])
    assert op_spec["import_row_count"] == 2
    assert "rows" not in op_spec["params"]

    receipt_row = project.db.execute(
        "SELECT * FROM receipts WHERE id=?", (result.receipt_id,)
    ).fetchone()
    assert receipt_row is not None
    assert receipt_row["action_kind"] == "import.csv"
    receipt = Receipt.model_validate(json.loads(receipt_row["body"]))
    assert receipt.schema_version == "frisket.receipt.v1"
    assert receipt.action_kind == "import.csv"
    assert receipt.idempotency_key == "import_csv@sha256:stable"
    assert receipt.op_ids == [1]
    file_read = next(
        item.ref for item in receipt.inputs if item.ref["kind"] == "local_file_read"
    )
    raw = seeded["csv_path"].read_bytes()
    assert file_read == {
        "kind": "local_file_read",
        "path": str(seeded["csv_path"]),
        "sha256": "sha256:" + hashlib.sha256(raw).hexdigest(),
        "byte_count": len(raw),
    }
    assert result.outputs[0].ref["columns"] == {
        name: column["id"] for name, column in columns.items()
    }
    assert any(item.ref["kind"] == "source_cell" for item in receipt.evidence)


CASES = [
    ExecutorCase(
        kind="import.csv",
        catalog=CatalogEntry(
            execution_mode="whole_project",
            async_mode="sync",
            writes_project=True,
            receipt_policy="writes_receipt",
            required_capabilities=("project:write",),
            side_effects=frozenset(
                {
                    "read_local_file",
                    "create_sheet",
                    "create_columns",
                    "create_rows",
                    "write_op",
                    "write_receipt",
                }
            ),
            error_codes=frozenset(
                {
                    "invalid_action_request",
                    "invalid_file_source",
                    "idempotency_conflict",
                }
            ),
            input_schema_properties=("sources", "columns"),
            output_schema_properties=(),
        ),
        seed=_seed,
        make_action=_make_action,
        gates=(
            Gate(
                "missing_idempotency_key",
                _missing_idempotency_action,
                "invalid_action_request",
            ),
            Gate(
                "invalid_file_source",
                _missing_source_action,
                "invalid_file_source",
            ),
            Gate(
                "csv_header_mismatch",
                _header_mismatch_action,
                "csv_header_mismatch",
            ),
            Gate("invalid_csv_value", _invalid_value_action, "invalid_csv_value"),
        ),
        expect_counts={"sheets": 1, "columns": 3, "rows": 2, "ops": 1, "receipts": 1},
        check_state=_check_state,
        request_style="typed",
    )
]


def test_import_csv_cli_round_trip(tmp_path: Path) -> None:
    """The `action validate` / `action run` CLI surface for a file-backed
    import: exit codes and parseable payloads (the shared harness runs
    in-process, so the CLI plumbing is only proven here)."""
    from frisket.actions.system import RootActionValidationResult
    from frisket.contracts.action import ActionResult

    csv_path = _write_csv(tmp_path, "rows.csv", CSV_BODY)
    spec_path = write_json(tmp_path, "import_csv.json", _import_csv_action(csv_path))

    cli_validate = run_cli("action", "validate", str(spec_path))
    assert cli_validate.returncode == 0, cli_validate.stderr
    validated = RootActionValidationResult.model_validate(
        json.loads(cli_validate.stdout)
    )
    assert validated.ok is True
    assert validated.action is not None
    assert validated.action.action_id == "import.csv"

    project_path = tmp_path / "csv.frisket"
    Project.create(project_path, name="CSV executor").close()
    run = run_cli(
        "action",
        "run",
        "--project",
        str(project_path),
        "--project-id",
        "project-csv",
        str(spec_path),
    )
    assert run.returncode == 0, run.stderr
    result = ActionResult.model_validate(json.loads(run.stdout))
    assert result.status == "completed"
    assert result.receipt_id is not None

    failing = _missing_source_action({"csv_path": csv_path, "dir": tmp_path})
    failing_path = write_json(tmp_path, "missing_source.json", failing)
    failed = run_cli(
        "action",
        "run",
        "--project",
        str(project_path),
        "--project-id",
        "project-csv",
        str(failing_path),
    )
    assert failed.returncode == 1
    failed_result = ActionResult.model_validate(json.loads(failed.stdout))
    assert failed_result.status == "failed"
    assert failed_result.errors[0].code == "invalid_file_source"


def test_import_csv_default_allows_more_than_10000_file_rows(tmp_path: Path) -> None:
    """File-backed Solo imports have no legacy total-row refusal either."""
    from frisket.engine.executor.actions import run_action_spec

    csv_path = _write_csv(
        tmp_path,
        "large.csv",
        "name\n" + "".join(f"Record {index}\n" for index in range(10_001)),
    )
    action = _import_csv_action(
        csv_path,
        sheet_name="Large CSV",
        idempotency_key="import_csv@sha256:10001",
    )
    action["params"]["columns"] = [{"name": "name", "type": "text"}]

    project = Project.create(tmp_path / "large-csv.frisket", name="Large CSV")
    try:
        result = run_action_spec(project, action, project_id="large-csv")

        assert result.status == "completed", result.errors
        sheet = project.db.execute(
            "SELECT id FROM sheets WHERE name='Large CSV'"
        ).fetchone()
        assert sheet is not None
        assert project.row_count(int(sheet["id"])) == 10_001
        receipt = project.db.execute(
            "SELECT body FROM receipts WHERE id=?", (result.receipt_id,)
        ).fetchone()
        assert receipt is not None
        body = str(receipt["body"])
        assert len(body.encode("utf-8")) < 100_000
        assert "Record 10000" not in body
        op = project.db.execute(
            "SELECT spec FROM ops WHERE id=?", (result.op_ids[0],)
        ).fetchone()
        assert op is not None
        assert json.loads(op["spec"])["import_row_count"] == 10_001
        assert "rows" not in json.loads(op["spec"])["params"]
    finally:
        project.close()


def test_import_csv_executes_a_real_source_larger_than_5mb(tmp_path: Path) -> None:
    """Typed streaming ingestion does not restore the retired source-byte cap."""
    from frisket.engine.executor.actions import run_action_spec

    payload = "x" * (5 * 1024 * 1024 + 1)
    csv_path = _write_csv(tmp_path, "over-5mb.csv", f"name\n{payload}\n")
    assert csv_path.stat().st_size > 5 * 1024 * 1024
    action = _import_csv_action(
        csv_path,
        sheet_name="Over 5 MB",
        idempotency_key="import_csv@sha256:over-5mb",
    )
    action["params"]["columns"] = [{"name": "name", "type": "text"}]

    project = Project.create(tmp_path / "over-5mb.frisket", name="Over 5 MB")
    try:
        result = run_action_spec(project, action, project_id="over-5mb")

        assert result.status == "completed", result.errors
        sheet = project.db.execute(
            "SELECT id FROM sheets WHERE name='Over 5 MB'"
        ).fetchone()
        assert sheet is not None
        assert project.row_count(int(sheet["id"])) == 1
    finally:
        project.close()


def test_csv_field_policy_is_configured_at_package_startup() -> None:
    """A new Frisket process lifts CPython's CSV field ceiling."""
    subprocess.run(
        [
            # subprocess-boundary: CSV startup policy needs a fresh interpreter
            sys.executable,
            "-c",
            "import csv; import frisket; "
            "assert csv.field_size_limit() == frisket.CSV_FIELD_SIZE_LIMIT",
        ],
        check=True,
    )


def test_csv_field_policy_falls_back_when_c_long_is_32_bit(monkeypatch: Any) -> None:
    """Windows accepts its C-long maximum rather than crashing at import."""
    import frisket

    requested: list[int] = []

    def windows_sized_limit(value: int) -> int:
        requested.append(value)
        if value > (1 << 31) - 1:
            raise OverflowError
        return value

    monkeypatch.setattr(frisket._csv, "field_size_limit", windows_sized_limit)

    assert frisket._configure_csv_field_size_limit() == (1 << 31) - 1
    assert requested == [sys.maxsize, (1 << 31) - 1]


def test_csv_field_policy_is_stable_across_concurrent_imports(
    tmp_path: Path, monkeypatch: Any
) -> None:
    """Concurrent requests retain the active CSV parser policy."""
    import csv
    from concurrent.futures import ThreadPoolExecutor

    from frisket.engine.executor import run_action_spec

    path = _write_csv(tmp_path, "large-fields.csv", f"name\n{'x' * 200_000}\n")
    initial = csv.field_size_limit()
    field_size_limit = csv.field_size_limit

    def reject_request_time_mutation(*args: int) -> int:
        assert not args, "CSV requests must not mutate process parser policy"
        return field_size_limit()

    monkeypatch.setattr(csv, "field_size_limit", reject_request_time_mutation)

    def import_one(index: int) -> int:
        project = Project.create(
            tmp_path / f"parallel-{index}.frisket", name="Parallel"
        )
        try:
            action = _import_csv_action(path)
            action["params"]["columns"] = [{"name": "name", "type": "text"}]
            result = run_action_spec(project, action, project_id="parallel")
            assert result.status == "completed", result.errors
            return project.row_count(result.outputs[0].sheet_id)
        finally:
            project.close()

    with ThreadPoolExecutor(max_workers=4) as pool:
        counts = list(pool.map(import_one, range(8)))

    assert counts == [1] * 8
    assert csv.field_size_limit() == initial


def test_typed_file_imports_accept_10001_rows_without_a_deployment_cap(
    tmp_path: Path,
) -> None:
    """The complete typed import path has no fixed total-row ceiling."""
    from openpyxl import Workbook

    from frisket.engine.executor import run_action_spec

    count = 10_001
    xlsx_path = tmp_path / "many.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet()
    worksheet.append(["name"])
    for index in range(count):
        worksheet.append([f"record-{index}"])
    workbook.save(xlsx_path)

    geojson_path = tmp_path / "many.geojson"
    geojson_path.write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "properties": {},
                        "geometry": {
                            "type": "Point",
                            "coordinates": [index, 0],
                        },
                    }
                    for index in range(count)
                ],
            }
        ),
        encoding="utf-8",
    )

    kml_path = tmp_path / "many.kml"
    kml_path.write_text(
        '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
        + "".join(
            "<Placemark><name>record-"
            f"{index}</name><Point><coordinates>{index},0</coordinates></Point>"
            "</Placemark>"
            for index in range(count)
        )
        + "</Document></kml>",
        encoding="utf-8",
    )

    one_text_column = [{"name": "name", "type": "text"}]
    project = Project.create(tmp_path / "many-project")
    try:
        for kind, source_path, options in (
            (
                "import.xlsx",
                xlsx_path,
                {"columns": one_text_column, "header": "present"},
            ),
            ("import.geojson", geojson_path, {"property_columns": []}),
            ("import.kml", kml_path, {"property_columns": []}),
        ):
            result = run_action_spec(
                project,
                {
                    "action_id": kind,
                    "scope": {"kind": "project"},
                    "sheet_name": f"Many {kind}",
                    "params": {
                        "source": {"kind": "file", "path": str(source_path)},
                        **options,
                    },
                    "idempotency_key": f"many-{kind}",
                },
                project_id="p",
            )
            assert result.status == "completed", result.errors
            assert result.outputs[0].ref["row_count"] == count
            assert project.row_count(result.outputs[0].ref["sheet_id"]) == count
    finally:
        project.close()

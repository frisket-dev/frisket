"""Workbook adapters that expose canonical typed cell text."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterator

from .models import RawCell, RawRow, SheetInfo, SourceKind
from .reader import TableReadError


def inspect_workbook(path: Path, *, kind: SourceKind) -> tuple[SheetInfo, ...]:
    if kind == "xlsx":
        from openpyxl import load_workbook

        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                return tuple(
                    SheetInfo(
                        index=index,
                        name=sheet.title,
                        hidden=sheet.sheet_state != "visible",
                        empty=_xlsx_sheet_empty(sheet),
                    )
                    for index, sheet in enumerate(workbook.worksheets)
                )
            finally:
                workbook.close()
    workbook = _open_calamine(path)
    try:
        return tuple(
            SheetInfo(
                index=index,
                name=metadata.name,
                hidden=str(metadata.visible).rsplit(".", 1)[-1] != "Visible",
                empty=workbook.get_sheet_by_index(index).start is None,
            )
            for index, metadata in enumerate(workbook.sheets_metadata)
        )
    finally:
        workbook.close()


class WorkbookRows:
    def __init__(self, path: Path, *, kind: SourceKind, worksheet: str | None) -> None:
        self._kind = kind
        self._source = None
        self._workbook: Any
        if kind == "xlsx":
            from openpyxl import load_workbook

            self._source = path.open("rb")
            self._workbook = load_workbook(self._source, read_only=True, data_only=True)
            try:
                sheet = (
                    self._workbook[worksheet]
                    if worksheet is not None
                    else self._workbook.active
                )
            except KeyError as exc:
                self.close()
                raise TableReadError(f"worksheet is unavailable: {worksheet}") from exc
            if sheet is None:
                self.close()
                raise TableReadError("workbook has no worksheet")
            self._rows = iter(sheet.iter_rows())
        else:
            self._workbook = _open_calamine(path)
            try:
                sheet = (
                    self._workbook.get_sheet_by_name(worksheet)
                    if worksheet is not None
                    else self._workbook.get_sheet_by_index(0)
                )
            except Exception as exc:
                self.close()
                raise TableReadError(f"worksheet is unavailable: {exc}") from exc
            self._rows = iter(sheet.iter_rows())
        self._ordinal = 0
        self._closed = False

    def __iter__(self) -> Iterator[RawRow]:
        return self

    def __next__(self) -> RawRow:
        try:
            values = next(self._rows)
        except StopIteration:
            raise
        except Exception as exc:
            raise TableReadError(f"workbook row could not be read: {exc}") from exc
        self._ordinal += 1
        cells = tuple(
            _xlsx_cell(value) if self._kind == "xlsx" else _typed_cell(value)
            for value in values
        )
        return RawRow(self._ordinal, self._ordinal, self._ordinal, cells)

    def close(self) -> None:
        workbook = getattr(self, "_workbook", None)
        if workbook is not None:
            workbook.close()
            self._workbook = None
        if self._source is not None:
            self._source.close()
            self._source = None
        self._closed = True


def _xlsx_sheet_empty(sheet: Any) -> bool:
    if sheet.max_row != 1 or sheet.max_column != 1:
        return False
    row = next(sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=1), ())
    return not row or row[0].value is None


def _xlsx_cell(cell: Any) -> RawCell:
    if getattr(cell, "data_type", None) == "e":
        return RawCell(str(cell.value or ""), "error")
    if isinstance(cell.value, datetime):
        from openpyxl.styles.numbers import is_datetime

        if is_datetime(cell.number_format) == "date":
            return RawCell(cell.value.date().isoformat(), "date")
    return _typed_cell(cell.value)


def _typed_cell(value: Any) -> RawCell:
    if value is None:
        return RawCell("", "blank")
    if isinstance(value, bool):
        return RawCell("true" if value else "false", "boolean")
    if isinstance(value, datetime):
        return RawCell(value.isoformat(), "datetime")
    if isinstance(value, date):
        return RawCell(value.isoformat(), "date")
    if isinstance(value, (int, float)):
        return RawCell(str(value), "number")
    if isinstance(value, (time, timedelta)):
        return RawCell(str(value), "text")
    return RawCell(str(value), "text")


def _open_calamine(path: Path) -> Any:
    try:
        from python_calamine import CalamineWorkbook
    except ImportError as exc:  # pragma: no cover - primary owns dependency setup.
        raise TableReadError("XLS and ODS reading requires python-calamine") from exc
    try:
        # The staged source may deliberately have no filename extension.
        # Calamine's object entry point detects the workbook from its bytes.
        return CalamineWorkbook.from_object(path)
    except Exception as exc:
        raise TableReadError(f"workbook could not be opened: {exc}") from exc

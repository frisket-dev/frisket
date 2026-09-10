"""Typed, lazy XLSX import using an admitted seekable source."""

from __future__ import annotations

from typing import Literal

from pydantic import Field, model_validator

from frisket.actions.core import ActionCategory, RowScope, action, create_sheet
from frisket.actions.imports import FileSource
from frisket.actions.types import (
    ActionParams,
    DynamicOutput,
    LocalFileReader,
    TableColumn,
    TableError,
    TableResult,
    TableRow,
    RowsAppender,
    AppendedRows,
    RowsUpdater,
    UpdatedImportedRows,
)
from frisket.contracts.action import MAX_IMPORT_ROWS_COLUMNS
from frisket.contracts.actions.schemas.imports import ValidatedImportColumns
from frisket.csv_values import parse_csv_value


class ImportXlsxParams(ActionParams):
    source: FileSource
    columns: ValidatedImportColumns = Field(
        min_length=1, max_length=MAX_IMPORT_ROWS_COLUMNS
    )
    header: Literal["present"] = "present"
    worksheet: str | None = None


class UpdateXlsxParams(ImportXlsxParams):
    key_columns: list[str] = Field(min_length=1)
    keep_existing_on_blank: bool = False

    @model_validator(mode="after")
    def _validate_update_mapping(self) -> UpdateXlsxParams:
        names = {column.name for column in self.columns}
        if len(self.key_columns) != len(set(self.key_columns)) or not set(
            self.key_columns
        ).issubset(names):
            raise ValueError("key_columns must uniquely name imported columns")
        if names == set(self.key_columns):
            raise ValueError("choose at least one non-key update column")
        return self


def _columns(params: ImportXlsxParams) -> tuple[TableColumn, ...]:
    return tuple(
        TableColumn(
            key=column.name,
            type=column.type,
            format=column.format,
            hidden=column.hidden,
        )
        for column in params.columns
    )


def import_xlsx(
    params: ImportXlsxParams, files: LocalFileReader
) -> TableResult[DynamicOutput]:
    def rows():
        source_names = [column.source_name or column.name for column in params.columns]
        try:
            from openpyxl import load_workbook

            with files.open_binary(params.source.path) as binary:
                formula_cells: set[tuple[int, int]] = set()
                if isinstance(params, UpdateXlsxParams):
                    formula_workbook = load_workbook(
                        binary, read_only=True, data_only=False
                    )
                    try:
                        formula_sheet = (
                            formula_workbook[params.worksheet]
                            if params.worksheet
                            else formula_workbook.active
                        )
                        if formula_sheet is None:
                            raise ValueError("workbook has no worksheet")
                        formula_rows = formula_sheet.iter_rows()
                        formula_header = next(formula_rows, None)
                        formula_names = [
                            str(cell.value) if cell.value is not None else f"col{index}"
                            for index, cell in enumerate(formula_header or [])
                        ]
                        selected_positions = {
                            formula_names.index(name)
                            for name in source_names
                            if formula_names.count(name) == 1
                        }
                        for row_index, cells in enumerate(formula_rows, start=2):
                            for position in selected_positions:
                                if (
                                    position < len(cells)
                                    and cells[position].data_type == "f"
                                ):
                                    formula_cells.add((row_index, position))
                    finally:
                        formula_workbook.close()
                    binary.seek(0)
                workbook = load_workbook(binary, read_only=True, data_only=True)
                try:
                    sheet = (
                        workbook[params.worksheet]
                        if params.worksheet
                        else workbook.active
                    )
                    if sheet is None:
                        raise ValueError("workbook has no worksheet")
                    worksheet_rows = sheet.iter_rows(values_only=True)
                    header = next(worksheet_rows, None)
                    actual = [
                        str(value) if value is not None else f"col{index}"
                        for index, value in enumerate(header or [])
                    ]
                    explicitly_mapped = any(
                        column.source_name is not None for column in params.columns
                    )
                    if (not explicitly_mapped and actual != source_names) or any(
                        actual.count(name) != 1 for name in source_names
                    ):
                        raise TableError(
                            "xlsx_header_mismatch",
                            "XLSX header must match declared column names",
                            details={"expected": source_names, "actual": actual},
                        )
                    positions = [actual.index(name) for name in source_names]
                    for row_index, values in enumerate(worksheet_rows, start=2):
                        unresolved_selected_formula = any(
                            (row_index, position) in formula_cells
                            for position in positions
                        )
                        if (
                            values is None
                            or all(value is None for value in values)
                            and not unresolved_selected_formula
                        ):
                            continue
                        row = {}
                        for column, index in zip(
                            params.columns, positions, strict=True
                        ):
                            value = values[index] if index < len(values) else None
                            if value is None and (row_index, index) in formula_cells:
                                raise TableError(
                                    "unresolved_xlsx_formula",
                                    "Selected XLSX formula has no cached result",
                                    details={
                                        "line": row_index,
                                        "column": column.name,
                                    },
                                )
                            text = "" if value is None else str(value)
                            try:
                                row[column.name] = parse_csv_value(text, column.type)
                            except ValueError as exc:
                                raise TableError(
                                    "invalid_xlsx_value",
                                    "XLSX value could not be coerced to declared column type",
                                    details={
                                        "line": row_index,
                                        "column": column.name,
                                        "type": column.type,
                                        "value": text,
                                    },
                                ) from exc
                        yield TableRow(output=DynamicOutput(row))
                finally:
                    workbook.close()
        except TableError:
            raise
        except Exception as exc:  # noqa: BLE001 - parser dependency/errors are data
            raise TableError(
                "xlsx_parse_failed",
                "XLSX workbook could not be parsed",
                details={"error": str(exc)},
            ) from exc

    return TableResult(
        rows=rows(),
        source={"kind": "file", "label": params.source.label, "importer": "xlsx"},
    )


XLSX = action(
    examples=(
        ImportXlsxParams(
            source=FileSource(kind="file", path="examples/data.xlsx"),
            columns=[{"name": "title", "type": "text"}],
            worksheet="Sheet1",
        ),
    ),
    name="xlsx",
    title="Import XLSX",
    description="Import a local XLSX worksheet using its declared column schema.",
    category=ActionCategory.CONVERT,
    run=create_sheet(import_xlsx, columns_from=_columns),
)


def append_xlsx(params: ImportXlsxParams, appender: RowsAppender) -> AppendedRows:
    return appender.append_xlsx(params)


APPEND_XLSX = action(
    examples=(
        ImportXlsxParams(
            source={"kind": "file", "path": "examples/data.xlsx"},
            columns=[{"name": "title", "type": "text"}],
        ),
    ),
    name="append_xlsx",
    title="Append XLSX rows",
    description="Stream one admitted XLSX mapping into an explicitly scoped sheet.",
    category=ActionCategory.CONVERT,
    run=append_xlsx,
    row_scope=RowScope.ALL_ROWS,
)


def update_xlsx(params: UpdateXlsxParams, updater: RowsUpdater) -> UpdatedImportedRows:
    return updater.update_xlsx(params)


UPDATE_XLSX = action(
    examples=(
        UpdateXlsxParams(
            source={"kind": "file", "path": "examples/data.xlsx"},
            columns=[{"name": "id", "type": "text"}, {"name": "title", "type": "text"}],
            key_columns=["id"],
        ),
    ),
    name="update_xlsx",
    title="Update rows from XLSX",
    description="Stream XLSX rows into exact-key updates on an explicitly scoped sheet.",
    category=ActionCategory.CONVERT,
    run=update_xlsx,
    row_scope=RowScope.ALL_ROWS,
)
